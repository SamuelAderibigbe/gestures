import datetime as dt

import pandas as pd
import pytest
from openpyxl import load_workbook

from gesture_pipeline.validation import match, parse_time, prepare, validate

from .conftest import needs_ffmpeg


@pytest.mark.parametrize("v,expected", [("25:12", 1512), ("25:12.5", 1512.5), ("1:02:03", 3723),
                                        ("75.5", 75.5), (90, 90), ("０:30".replace("０", "0"), 30),
                                        (dt.time(0, 25, 12), 1512), ("", None), (None, None)])
def test_parse_time(v, expected):
    assert parse_time(v) == expected


@pytest.mark.parametrize("v,expected", [(3, 3), (3.0, 3), ("3", 3), ("P3", 3), ("p 12", 12), ("3.0", 3),
                                        ("abc", None), (float("nan"), None)])
def test_parse_person(v, expected):
    from gesture_pipeline.validation import parse_person
    assert parse_person(v) == expected


def test_match_is_one_to_one_with_tolerance():
    a = pd.DataFrame([{"person_id": 1, "gesture": "nod", "start_s": 10, "end_s": 11},
                      {"person_id": 1, "gesture": "nod", "start_s": 20, "end_s": 21}])
    b = pd.DataFrame([{"person_id": 1, "gesture": "nod", "start_s": 11.5, "end_s": 12},   # within 1 s
                      {"person_id": 2, "gesture": "nod", "start_s": 20, "end_s": 21}])    # wrong person
    tp, miss_a, miss_b = match(a, b, tol=1.0)
    assert tp == 1 and miss_a == [1] and miss_b == [1]


@needs_ffmpeg
def test_prepare_and_validate_on_demo(tmp_path):
    from gesture_pipeline.demo import run_demo
    out = tmp_path / "demo"
    run_demo(out)
    val = tmp_path / "validation"
    clips = prepare(out, val, n_clips=1, length_s=39, seed=0)       # (almost) the whole 40 s demo
    assert (val / "clips" / "C01.mp4").exists() and (val / "INSTRUCTIONS.md").exists()

    # A "coder" marks the scripted gestures, but misses one nod and adds a nod nobody scripted.
    wb = load_workbook(val / "coding_sheet.xlsx")
    ws = wb["annotations"]
    pid = pd.read_csv(out / "demo_conference" / "people.csv").set_index("role_key")
    panel = sorted(pd.read_csv(out / "demo_conference" / "people.csv").query("role_key == 'panel'").person_id)
    aud = sorted(pd.read_csv(out / "demo_conference" / "people.csv").query("role_key == 'audience'").person_id)
    start = float(clips.iloc[0].start_s)
    assert start < 4.5 and clips.iloc[0].end_s > 28.5                # clip covers both scripted nods
    ws.append(["C01", panel[0], "nod", "00:05", "00:06.5"])          # scripted nod #1 (nod #2 at 0:27 missed)
    ws.append(["C01", panel[2], "nod", "00:15", "00:16"])            # not scripted -> program "misses" it
    ws.append(["C01", aud[0], "hand_raise", "0:10", "0:13"])
    cg = wb["coded_gestures"]
    for row in cg.iter_rows(min_row=2):
        row[1].value = "yes" if row[0].value in ("nod", "hand_raise") else "no"
    wb.save(val / "coding_sheet_tester.xlsx")

    acc = validate(out, val)["Accuracy"].set_index("gesture")
    assert acc.loc["hand_raise", "precision"] == 1.0 and acc.loc["hand_raise", "recall"] == 1.0
    assert acc.loc["nod", "found_by_program"] == 2 and acc.loc["nod", "marked_by_person"] == 2
    assert acc.loc["nod", "matched"] == 1                              # one agreed, one each way missed
    assert "holding_mic" not in acc.index                              # not coded -> not counted
